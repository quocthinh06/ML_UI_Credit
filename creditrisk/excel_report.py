"""Excel reports built from the exact exported rows, with editable native charts."""

import numpy as np
import pandas as pd
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.data_source import NumData, NumVal, StrData, StrRef, StrVal
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.metrics import confusion_matrix, roc_curve

from .analytics import evaluation, labelled, summary
from .validation import TARGET


def write_report(frame, path):
    stats = summary(frame)
    insights = [
        ("Nội dung", "Phân tích rủi ro tín dụng trên đúng các hồ sơ đang xuất"),
        ("Số dòng dữ liệu (không gồm tiêu đề)", len(frame)),
        ("Số cột dữ liệu", len(frame.columns)),
        ("Các cột", ", ".join(map(str, frame.columns))),
    ]
    for column in (
        "model_name", "model_version", "model_sha256", "dataset_version_id", "run_id",
        "decision_threshold", "decision_basis", "risk_low_threshold", "risk_high_threshold",
    ):
        if column in frame:
            insights.append((column, ", ".join(map(str, frame[column].dropna().unique()))))
    insights.extend([
        ("PD trung bình", stats["mean_pd"]),
        ("Số hồ sơ cảnh báo", stats["warnings"]),
        ("Số hồ sơ rủi ro cao", stats["high_risk"]),
        ("Số nhãn thực tế hợp lệ", stats["labelled"]),
        ("Tỷ lệ vỡ nợ thực tế", stats["actual_default_rate"]),
    ])
    metrics = evaluation(frame)
    insights.extend(metrics.items())
    if stats["mean_pd"] is not None:
        insights.append(("Nhận định", (
            f"{stats['warnings']:,}/{len(frame):,} hồ sơ được cảnh báo; "
            f"{stats['high_risk']:,} hồ sơ rủi ro cao. PD trung bình {stats['mean_pd']:.2%}."
        )))
    insights.append(("Phạm vi đánh giá", (
        "Chỉ tính chỉ số trên nhãn thực tế 0/1 hợp lệ; AUC cần cả hai lớp."
        if metrics else "Chưa đủ nhãn thực tế và kết quả dự báo để đánh giá model."
    )))

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Keep raw rows first for existing consumers that read the first worksheet.
        frame.to_excel(writer, sheet_name="Du lieu", index=False)
        pd.DataFrame(insights, columns=["Nội dung / chỉ số", "Giá trị"]).to_excel(
            writer, sheet_name="Insight", index=False, na_rep="Không có dữ liệu"
        )
        book = writer.book
        dashboard = book.create_sheet("Dashboard")
        dashboard.append(["DASHBOARD · " + (
            ", ".join(map(str, frame.model_name.dropna().unique()))
            if "model_name" in frame else "Dữ liệu chưa dự báo"
        )])
        dashboard.append([f"Phạm vi: {len(frame):,} dòng · {len(frame.columns):,} cột"])
        source = book.create_sheet("So lieu bieu do")
        next_row = 1
        chart_index = 0

        def chart(title, labels, values, *, percent=False, category_title="Nhóm"):
            nonlocal next_row, chart_index
            labels, values = list(map(str, labels)), list(values)
            start = next_row
            source.cell(start, 1, title)
            source.cell(start, 2, "Giá trị")
            for row, (label, value) in enumerate(zip(labels, values), start + 1):
                source.cell(row, 1, str(label))
                source.cell(row, 2, None if pd.isna(value) else float(value))
                if percent:
                    source.cell(row, 2).number_format = "0.0%"
            end = start + len(labels)
            if end == start:
                return
            plot = BarChart()
            plot.title = title
            plot.add_data(Reference(source, min_col=2, min_row=start, max_row=end), titles_from_data=True)
            plot.set_categories(Reference(source, min_col=1, min_row=start + 1, max_row=end))
            # Cache labels too: spreadsheet viewers may not resolve sheet references on opening.
            series = plot.series[0]
            series.cat.numRef = None
            series.cat.strRef = StrRef(
                f=f"'So lieu bieu do'!$A${start + 1}:$A${end}",
                strCache=StrData(ptCount=len(labels), pt=[StrVal(idx=i, v=v) for i, v in enumerate(labels)]),
            )
            series.val.numRef.numCache = NumData(
                formatCode="0.0%" if percent else "#,##0",
                ptCount=len(values),
                pt=[NumVal(idx=i, v=float(v)) for i, v in enumerate(values) if not pd.isna(v)],
            )
            plot.x_axis.title = category_title
            plot.y_axis.title = "Tỷ lệ (%)" if percent else "Số hồ sơ"
            plot.x_axis.delete = plot.y_axis.delete = False
            plot.x_axis.tickLblPos = "low"
            plot.y_axis.tickLblPos = "nextTo"
            plot.x_axis.tickLblSkip = 1
            plot.dLbls = DataLabelList(
                showVal=True, showCatName=False, showSerName=False, showLegendKey=False,
                dLblPos="outEnd", numFmt="0.0%" if percent else "#,##0",
            )
            plot.legend = None
            plot.width, plot.height = 22, 11
            if percent:
                plot.y_axis.numFmt = "0%"
            dashboard.add_chart(plot, f"{'A' if chart_index % 2 == 0 else 'M'}{4 + chart_index // 2 * 23}")
            next_row = end + 3
            chart_index += 1

        if "probability" in frame:
            counts = frame.risk_level.value_counts().reindex(["Low", "Medium", "High"], fill_value=0)
            chart("Phân nhóm rủi ro từ PD", ["Thấp", "Trung bình", "Cao"], counts.values,
                  category_title="Mức rủi ro (risk_level)")
            counts, edges = np.histogram(frame.probability.dropna(), bins=np.linspace(0, 1, 21))
            chart("Phân bố PD đã hiệu chỉnh", [f"{a:.0%}–{b:.0%}" for a, b in zip(edges[:-1], edges[1:])],
                  counts, category_title="Khoảng xác suất PD (probability)")
            if "AGE" in frame:
                groups = pd.cut(frame.AGE, [17, 29, 39, 49, 59, 100],
                                labels=["18–29", "30–39", "40–49", "50–59", "60+"])
                means = frame.groupby(groups, observed=False).probability.mean()
                chart("PD trung bình theo tuổi", means.index, means.values, percent=True,
                      category_title="Nhóm tuổi (AGE)")
            counts = frame.default_label.value_counts().reindex([0, 1], fill_value=0)
            chart("Nhãn cảnh báo theo ngưỡng đã lưu", ["Nhãn 0", "Nhãn 1 · cảnh báo"], counts.values,
                  category_title="Nhãn dự báo (default_label)")
        for column, title in [("AGE", "Phân bố tuổi"), ("LIMIT_BAL", "Phân bố hạn mức (TWD)"),
                              ("EDUCATION", "Trình độ học vấn"), ("PAY_0", "Trạng thái trả nợ gần nhất")]:
            if column in frame:
                values = frame[column].dropna()
                if column in ("AGE", "LIMIT_BAL") and not values.empty:
                    counts, edges = np.histogram(values, bins=10)
                    chart(title, [f"{a:,.0f}–{b:,.0f}" for a, b in zip(edges[:-1], edges[1:])], counts,
                          category_title=column)
                else:
                    counts = values.value_counts().sort_index()
                    chart(title, counts.index, counts.values, category_title=column)
        if metrics:
            known = labelled(frame)
            cm = confusion_matrix(known[TARGET], known.default_label, labels=[0, 1])
            chart("Ma trận nhầm lẫn · thực tế / dự báo", [
                "TN · Thực tế 0 / Dự báo 0", "FP · Thực tế 0 / Dự báo 1",
                "FN · Thực tế 1 / Dự báo 0", "TP · Thực tế 1 / Dự báo 1",
            ], cm.ravel(), category_title="Nhãn thực tế / nhãn dự báo (0: không vỡ nợ, 1: vỡ nợ)")
            keys = ["Accuracy", "Precision", "Recall", "F1"]
            chart("Đánh giá model trên nhãn hợp lệ", keys, [metrics[k] for k in keys], percent=True,
                  category_title="Chỉ số đánh giá")
            counts = known[TARGET].value_counts().reindex([0, 1], fill_value=0)
            chart("Nhãn thực tế", ["Không vỡ nợ", "Vỡ nợ"], counts.values, category_title=TARGET)
            if known[TARGET].nunique() == 2:
                fpr, tpr, _ = roc_curve(known[TARGET], known.raw_score)
                start = next_row
                source.cell(start, 1, "ROC · FPR")
                source.cell(start, 2, "ROC · TPR")
                for row, (x, y) in enumerate(zip(fpr, tpr), start + 1):
                    source.cell(row, 1, float(x))
                    source.cell(row, 2, float(y))
                plot = ScatterChart()
                plot.title = "ROC · điểm thô"
                plot.x_axis.title, plot.y_axis.title = "False positive rate", "True positive rate"
                plot.x_axis.scaling.min = plot.y_axis.scaling.min = 0
                plot.x_axis.scaling.max = plot.y_axis.scaling.max = 1
                plot.series.append(Series(
                    Reference(source, min_col=2, min_row=start + 1, max_row=start + len(fpr)),
                    Reference(source, min_col=1, min_row=start + 1, max_row=start + len(fpr)),
                    title="ROC",
                ))
                plot.width, plot.height = 22, 11
                dashboard.add_chart(plot, f"{'A' if chart_index % 2 == 0 else 'M'}{4 + chart_index // 2 * 23}")
        for sheet in book:
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="087F83")
            if sheet != dashboard:
                sheet.auto_filter.ref = sheet.dimensions
                for cells in sheet.columns:
                    width = min(48, max(14, max(len(str(c.value or "")) for c in cells[:100]) + 2))
                    sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
            else:
                for col in range(1, 25):
                    sheet.column_dimensions[get_column_letter(col)].width = 10
        book["Insight"].column_dimensions["A"].width = 42
        book["Insight"].column_dimensions["B"].width = 100
        for row in book["Insight"].iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            book["Insight"].row_dimensions[row[0].row].height = 45
