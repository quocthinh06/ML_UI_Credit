import pandas as pd
import pytest
from openpyxl import load_workbook

from creditrisk.service import export_filename, export_frame
from creditrisk.validation import TARGET, Thresholds, ValidationError


def test_excel_report_uses_selected_model_and_exported_rows(frame, model, tmp_path):
    frame[TARGET] = 1
    frame["ID"] = "=1+1"
    path = tmp_path / "report.xlsx"
    export_frame(frame, path, model=model, thresholds=Thresholds())
    book = load_workbook(path)
    assert book.sheetnames == ["Du lieu", "Insight", "Dashboard", "So lieu bieu do"]
    data = pd.read_excel(path)
    assert len(data) == 1 and data.model_name.iloc[0] == model.name
    assert data.ID.iloc[0] == "'=1+1"
    insights = dict(book["Insight"].values)
    assert insights["Số dòng dữ liệu (không gồm tiêu đề)"] == len(data)
    assert insights["Số cột dữ liệu"] == len(data.columns)
    assert insights["PD trung bình"] == 0.25
    assert insights["model_sha256"] == model.digest
    assert insights["labelled_rows"] == 1
    assert len(book["Dashboard"]._charts) == 11
    for chart in book["Dashboard"]._charts:
        assert chart.x_axis.delete is False and chart.y_axis.delete is False
        assert chart.x_axis.tickLblPos == "low"
        assert chart.dLbls.showVal is True
        assert chart.dLbls.dLblPos == "outEnd"
        labels = chart.series[0].cat.strRef.strCache
        assert labels.ptCount == len(labels.pt) > 0
        assert chart.series[0].val.numRef.numCache.ptCount == labels.ptCount
    confusion = book["Dashboard"]._charts[8]
    assert [p.v for p in confusion.series[0].cat.strRef.strCache.pt] == [
        "TN · Thực tế 0 / Dự báo 0", "FP · Thực tế 0 / Dự báo 1",
        "FN · Thực tế 1 / Dự báo 0", "TP · Thực tế 1 / Dự báo 1",
    ]
    assert book["Dashboard"]._charts[2].dLbls.numFmt == "0.0%"
    assert "probability" not in frame


def test_excel_report_rejects_other_model_without_replacing_file(frame, model, tmp_path):
    path = tmp_path / "report.xlsx"
    path.write_bytes(b"existing")
    result = pd.concat([frame, model.predict(frame, Thresholds())], axis=1)
    result["model_sha256"] = "another model"
    with pytest.raises(ValidationError, match="model"):
        export_frame(result, path, model=model)
    assert path.read_bytes() == b"existing"


def test_excel_report_without_labels(frame, model, tmp_path):
    path = tmp_path / "report.xlsx"
    export_frame(frame, path, model=model)
    book = load_workbook(path)
    insights = dict(book["Insight"].values)
    assert insights["Số nhãn thực tế hợp lệ"] == 0
    assert "Accuracy" not in insights
    assert len(book["Dashboard"]._charts) == 8


def test_excel_report_preserves_filtered_results_and_roc(frame, model, tmp_path):
    rows = pd.concat([frame] * 3, ignore_index=True)
    rows[TARGET] = [0, 1, 1]
    result = pd.concat([rows, model.predict(rows, Thresholds())], axis=1)
    result["model_sha256"] = model.digest
    result["model_name"] = model.name
    result.loc[0, "probability"] = 0.1
    filtered = result.iloc[:2]
    path = tmp_path / "filtered.xlsx"
    export_frame(filtered, path, model=model)
    book = load_workbook(path)
    insights = dict(book["Insight"].values)
    assert insights["Số dòng dữ liệu (không gồm tiêu đề)"] == 2
    assert insights["PD trung bình"] == pytest.approx(0.175)
    assert insights["AUC"] == 0.5
    assert len(book["Dashboard"]._charts) == 12
    assert pd.read_excel(path).probability.tolist() == [0.1, 0.25]


@pytest.mark.parametrize(("model_name", "filename"), [
    ("Random Forest", "ket_qua_RF.xlsx"),
    ("Logistic Regression", "ket_qua_LR.xlsx"),
    ("XGBoost", "ket_qua_XGb.xlsx"),
    (None, "ket_qua.xlsx"),
])
def test_export_filename_for_selected_model(model_name, filename):
    assert export_filename(model_name) == filename
