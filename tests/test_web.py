from io import BytesIO
import gzip
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from creditrisk.web import create_app
from creditrisk.config import ROOT


@pytest.fixture
def client(monkeypatch, model):
    model.metadata = {"cv_auc": 0.8, "test_auc": 0.75}
    monkeypatch.setattr("creditrisk.web.model_registry", lambda: {model.name: model})
    return create_app({"TESTING": True}).test_client()


def test_web_entrypoint_and_assets(client):
    response = client.get("/")
    assert response.status_code == 200
    for page in ["overview", "datasets", "predict", "history", "models"]:
        assert f'id="page-{page}"'.encode() in response.data
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/app.css").status_code == 200
    assert "frame-ancestors 'none'" in response.headers["Content-Security-Policy"]
    assert client.get("/api/health").json["storage"] == "browser"


def test_web_models(client, model):
    data = client.get("/api/models").json
    assert data["models"][0]["name"] == model.name
    assert len(data["features"]) == 23


def test_login_configuration_is_explicit_without_database(client):
    response = client.get("/api/auth/session")
    assert response.json["user"] is None and not response.json["configured"]
    assert client.get("/login").status_code == 200
    assert client.post("/api/auth/login", json={"username": "not_configured", "password": "not-a-real-password"}).status_code == 503
    assert client.post("/api/auth/logout", json={}).status_code == 200


def test_web_prediction_uses_real_input_contract(client, frame, model):
    response = client.post("/api/predict", json={"rows": frame.to_dict("records"), "model": model.name})
    assert response.status_code == 200
    assert response.json["rows"][0]["probability"] == 0.25
    assert response.json["rows"][0]["model_sha256"] == model.digest
    assert response.json["thresholds"]["basis"] == "raw"


@pytest.mark.parametrize("body", [None, [], {}, {"rows": []}, {"rows": [{"AGE": 1}]},
                                     {"rows": ["not an object"]}])
def test_web_rejects_invalid_inputs(client, body):
    response = client.post("/api/predict", data=json.dumps(body), content_type="application/json")
    assert 400 <= response.status_code < 500


def test_web_batch_limit_and_unknown_model(client, frame):
    rows = frame.to_dict("records")
    assert client.post("/api/predict", json={"rows": rows * 501}).status_code == 400
    assert client.post("/api/predict", json={"rows": rows, "model": "missing"}).status_code == 400


def test_web_import_keeps_ids_and_labels(client, frame):
    frame["ID"] = "00042"
    frame["default.payment.next.month"] = 1
    response = client.post("/api/import", data={"file": (BytesIO(frame.to_csv(index=False).encode()), "data.csv")})
    assert response.status_code == 200
    assert response.json["rows"][0]["ID"] == "00042"
    assert response.json["rows"][0]["default.payment.next.month"] == 1
    assert client.post("/api/import", data={"file": (BytesIO(b"bad"), "model.pkl")}).status_code == 400


def test_web_export_recomputes_untrusted_scores(client, frame, model):
    rows = frame.to_dict("records")
    rows[0].update(probability=0.99, model_sha256=model.digest, run_id="run-1")
    body = gzip.compress(json.dumps({"rows": rows, "model": model.name}).encode())
    response = client.post("/api/export", data=body, content_type="application/json",
                           headers={"Content-Encoding": "gzip"})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.data))
    assert book.sheetnames == ["Du lieu", "Insight", "Dashboard", "So lieu bieu do"]
    output = pd.read_excel(BytesIO(response.data))
    assert output.probability.iloc[0] == 0.25
    assert output.run_id.iloc[0] == "run-1"
    assert "attachment" in response.headers["Content-Disposition"]


def test_web_export_rejects_other_model(client, frame, model):
    rows = frame.to_dict("records")
    rows[0]["model_sha256"] = "other-model"
    response = client.post("/api/export", json={"rows": rows, "model": model.name})
    assert response.status_code == 400


def test_web_json_compression(client):
    response = client.get("/api/sample", headers={"Accept-Encoding": "gzip"})
    assert response.headers["Content-Encoding"] == "gzip"
    assert len(json.loads(gzip.decompress(response.data))["rows"]) == 3


def test_web_cloud_workspace_requires_private_key(client, monkeypatch):
    monkeypatch.delenv("MONGODB_URI", raising=False)
    assert client.get("/api/workspace").status_code == 401
    assert client.get("/api/workspace", headers={"Authorization": "Bearer short"}).status_code == 401
    assert client.get("/api/workspace", headers={"Authorization": "Bearer " + "a" * 64}).status_code == 503


def test_web_cloud_workspace_is_scoped_to_hashed_key(client, monkeypatch):
    from hashlib import sha256

    seen = []
    def write(owner, state, revision):
        seen.append((owner, state, revision))
        return {"revision": revision + 1}
    monkeypatch.setattr("creditrisk.cloud_storage.write_snapshot", write)
    response = client.put("/api/workspace", headers={"Authorization": "Bearer " + "b" * 64},
                          json={"state": {"datasets": [], "runs": []}, "revision": 0})
    assert response.status_code == 200 and response.json["revision"] == 1
    assert seen[0][0] == sha256(("b" * 64).encode()).hexdigest()


@pytest.mark.skipif(not (ROOT / "models/model_xgb_v1.pkl").exists(), reason="Download trusted model handoffs")
def test_web_real_models_predict_and_explain():
    from creditrisk.service import read_file

    client = create_app({"TESTING": True}).test_client()
    rows = read_file(ROOT / "examples/sample_customers.csv").to_dict("records")
    names = [m["name"] for m in client.get("/api/models").json["models"]]
    assert set(names) == {"Random Forest", "Logistic Regression", "XGBoost"}
    probabilities = []
    for name in names:
        result = client.post("/api/predict", json={"rows": rows, "model": name})
        assert result.status_code == 200
        probabilities.append(result.json["rows"][0]["probability"])
        explanation = client.post("/api/explain", json={"rows": rows[:1], "model": name})
        assert explanation.status_code == 200
        assert len(explanation.json["rows"]) == 12
    assert len(set(probabilities)) == 3
